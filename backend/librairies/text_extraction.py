"""
librairies/text_extraction.py
==============================

Extraction "universelle" de texte a partir d'une piece jointe (mission RAG
§3 : "analyse universelle des pieces jointes"). Point d'entree unique :
`extract_text(raw, filename, mime_type)`.

Contrat herite de l'ancien `jobs._extract_attachment_text` (Phase 1, ne pas
affaiblir) : ne leve jamais, n'invente jamais un contenu. Un format non
reconnu ou une extraction qui echoue renvoie `(None, "message explicite")`,
jamais un texte fabrique. Utilise a la fois par :
- jobs.py (contexte du message envoye a ChatGPT/Claude, piece jointe par
  piece jointe) ;
- librairies/rag.py (ingestion RAG, meme extraction reutilisee telle quelle
  plutot que dupliquee -- mission RAG §4b : "reutiliser Phase 1").
"""

from __future__ import annotations

import io

# Extensions/types traites comme du texte brut (decodage direct, aucune
# dependance necessaire) -- liste ouverte de formats lisibles tels quels,
# jamais une simulation d'extraction pour un format non reconnu.
PLAIN_TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "csv", "tsv", "json", "log", "py", "js", "ts",
    "html", "htm", "css", "yaml", "yml", "xml", "ini", "cfg",
    "rst", "sh", "sql",
}

# Formats acceptes a l'upload (workspace_routes.ALLOWED_MIME_TYPES) meme
# quand le navigateur envoie `application/octet-stream` (frequent pour les
# fichiers de code) : source unique de verite partagee avec les routes, pour
# ne jamais laisser deriver "ce qui est accepte a l'upload" et "ce que cette
# extraction sait vraiment lire" (mission RAG §3).
KNOWN_TEXT_EXTENSIONS = PLAIN_TEXT_EXTENSIONS

_XLSX_ROW_CAP = 500


def _extension_of(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _extract_pdf(raw: bytes) -> tuple[str | None, str | None]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc:
        return None, f"extraction PDF impossible ({str(exc)[:150]})"
    if not text:
        return None, "PDF sans texte extractible (probablement une image scannee)"
    return text, None


def _extract_docx(raw: bytes) -> tuple[str | None, str | None]:
    try:
        import docx

        document = docx.Document(io.BytesIO(raw))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
    except Exception as exc:
        return None, f"extraction DOCX impossible ({str(exc)[:150]})"
    if not text:
        return None, "DOCX sans texte extractible"
    return text, None


def _extract_pptx(raw: bytes) -> tuple[str | None, str | None]:
    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(raw))
        parts: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            slide_parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    slide_parts.append(shape.text_frame.text.strip())
            if slide_parts:
                parts.append(f"--- Slide {index} ---\n" + "\n".join(slide_parts))
        text = "\n\n".join(parts).strip()
    except Exception as exc:
        return None, f"extraction PPTX impossible ({str(exc)[:150]})"
    if not text:
        return None, "PPTX sans texte extractible"
    return text, None


def _render_sheet_rows(rows_iter, row_cap: int) -> tuple[str, bool]:
    lines: list[str] = []
    truncated = False
    for i, row in enumerate(rows_iter):
        if i >= row_cap:
            truncated = True
            break
        values = ["" if cell is None else str(cell) for cell in row]
        if any(v.strip() for v in values):
            lines.append(" | ".join(values))
    return "\n".join(lines), truncated


def _extract_xlsx(raw: bytes) -> tuple[str | None, str | None]:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in workbook.worksheets:
            body, truncated = _render_sheet_rows(sheet.iter_rows(values_only=True), _XLSX_ROW_CAP)
            if not body:
                continue
            header = f"--- Sheet: {sheet.title} ---"
            if truncated:
                body += f"\n... ({_XLSX_ROW_CAP}+ lignes, affichage tronque)"
            parts.append(f"{header}\n{body}")
        text = "\n\n".join(parts).strip()
    except Exception as exc:
        return None, f"extraction XLSX impossible ({str(exc)[:150]})"
    if not text:
        return None, "classeur XLSX vide ou sans donnees lisibles"
    return text, None


def _extract_xls(raw: bytes) -> tuple[str | None, str | None]:
    try:
        import xlrd

        book = xlrd.open_workbook(file_contents=raw)
        parts: list[str] = []
        for sheet in book.sheets():
            rows = (sheet.row_values(r) for r in range(sheet.nrows))
            body, truncated = _render_sheet_rows(rows, _XLSX_ROW_CAP)
            if not body:
                continue
            header = f"--- Sheet: {sheet.name} ---"
            if truncated:
                body += f"\n... ({_XLSX_ROW_CAP}+ lignes, affichage tronque)"
            parts.append(f"{header}\n{body}")
        text = "\n\n".join(parts).strip()
    except Exception as exc:
        return None, f"extraction XLS impossible ({str(exc)[:150]})"
    if not text:
        return None, "classeur XLS vide ou sans donnees lisibles"
    return text, None


_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLS_MIME = "application/vnd.ms-excel"


def extract_text(raw: bytes, filename: str, mime_type: str) -> tuple[str | None, str | None]:
    """Renvoie (texte, None) en cas de succes, ou (None, "raison explicite")
    sinon. Ne fabrique jamais de contenu pour un format non pris en charge."""
    mime_type = (mime_type or "").lower()
    extension = _extension_of(filename or "")

    if mime_type.startswith("text/") or extension in PLAIN_TEXT_EXTENSIONS:
        return raw.decode("utf-8", errors="replace"), None

    if mime_type == "application/pdf" or extension == "pdf":
        return _extract_pdf(raw)

    if mime_type == _DOCX_MIME or extension == "docx":
        return _extract_docx(raw)

    if mime_type == _PPTX_MIME or extension == "pptx":
        return _extract_pptx(raw)

    if mime_type == _XLSX_MIME or extension in ("xlsx", "xlsm"):
        return _extract_xlsx(raw)

    if mime_type == _XLS_MIME or extension == "xls":
        return _extract_xls(raw)

    if extension == "doc" or mime_type == "application/msword":
        # Format binaire legacy : aucun lecteur pur-Python fiable sans
        # dependance externe (LibreOffice, etc.). On prefere un message
        # honnete a un contournement fragile qui pretendrait le supporter.
        return None, "format DOC (legacy binaire) non pris en charge pour l'extraction automatique"

    if mime_type.startswith("image/"):
        return None, "images : analyse automatique du contenu non disponible (pas de vision configuree cote n8n)"

    return None, f"type de fichier non pris en charge pour l'aperçu automatique ({extension or mime_type or 'inconnu'})"

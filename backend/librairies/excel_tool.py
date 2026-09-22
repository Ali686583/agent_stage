"""
librairies/excel_tool.py
==========================

Moteur DETERMINISTE de lecture/modification de classeurs Excel (mission
Excel, Phase 5). L'IA decide QUOI changer (via l'outil MCP edit_excel, voir
excel_tool_server.py) ; ce module decide COMMENT -- jamais l'inverse, l'IA
ne fabrique jamais elle-meme des octets XLSX.

Regle absolue (mission §Original preservation) : JAMAIS d'ecriture sur le
fichier source. Toute modification produit un NOUVEAU fichier (nouvelle
ligne `files`, nouveau `storage_reference`) ; l'original reste inchange et
reste telechargeable par son propre id.

Ce module tourne dans le process WEB (les routes /mcp/excel-edit sont des
routes Flask normales, appelees par n8n pendant l'execution de l'Agent) --
contrairement au worker RQ, ce process a acces au volume de stockage
(UPLOAD_DIR), donc lecture/ecriture directe sur disque ici, pas d'URL
signee (voir jobs.py/rag.py pour le contraste, worker sans acces disque).
"""

from __future__ import annotations

import os
import re
import uuid

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from librairies import workspace

UPLOAD_DIR = os.environ.get("UPLOAD_DIR", "/data/uploads")

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUPPORTED_OPERATIONS = {
    "set_cell", "set_range", "add_column", "delete_column", "add_row",
    "delete_row", "add_sheet", "rename_sheet", "delete_sheet",
    "write_formula", "copy_range", "sort_range", "filter_rows", "clear_range",
}


class ExcelToolError(RuntimeError):
    """Erreur explicite (mission §9) : jamais de succes invente. Le message
    est renvoye tel quel a l'outil MCP (isError: true)."""


def _storage_path(file_id: str) -> str:
    storage_reference = workspace.get_file_storage_reference(file_id)
    if not storage_reference:
        raise ExcelToolError(f"fichier introuvable ({file_id})")
    return os.path.join(UPLOAD_DIR, storage_reference)


def describe_workbook(file_id: str) -> dict:
    """Structure reelle du classeur (mission §5.1) : permet a l'Agent de
    fonder son edit_excel sur des noms de feuille/colonnes reels plutot que
    devines depuis le seul texte extrait (text_extraction.py, qui tronque a
    500 lignes)."""
    path = _storage_path(file_id)
    try:
        workbook = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        raise ExcelToolError(f"classeur illisible ({str(exc)[:150]})") from exc
    sheets = []
    for sheet in workbook.worksheets:
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())]
        sheets.append({
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "header": header,
        })
    return {"fileId": file_id, "sheets": sheets}


def _resolve_sheet(workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ExcelToolError(f"feuille introuvable ({sheet_name})")
        return workbook[sheet_name]
    return workbook.active


def _parse_cell_ref(ref: str) -> tuple[int, int]:
    match = re.match(r"^([A-Za-z]+)(\d+)$", (ref or "").strip())
    if not match:
        raise ExcelToolError(f"reference de cellule invalide ({ref})")
    return column_index_from_string(match.group(1).upper()), int(match.group(2))


def _parse_range_ref(ref: str) -> tuple[int, int, int, int]:
    parts = (ref or "").strip().split(":")
    if len(parts) != 2:
        raise ExcelToolError(f"plage invalide ({ref})")
    start_col, start_row = _parse_cell_ref(parts[0])
    end_col, end_row = _parse_cell_ref(parts[1])
    return start_col, start_row, end_col, end_row


def _apply_op_add_column(sheet, op: dict) -> str:
    header = str(op.get("header") or "").strip()
    if not header:
        raise ExcelToolError("add_column requiert 'header'")
    formula_template = op.get("formula")
    start_row = int(op.get("startRow") or 2)
    target_col = sheet.max_column + 1
    col_letter = get_column_letter(target_col)
    sheet.cell(row=1, column=target_col, value=header)
    changed = 0
    if formula_template:
        for row in range(start_row, sheet.max_row + 1):
            if all(sheet.cell(row=row, column=c).value is None for c in range(1, target_col)):
                continue
            formula = formula_template.replace("{row}", str(row))
            sheet.cell(row=row, column=target_col, value=formula)
            changed += 1
    return f"colonne '{header}' ajoutee en {col_letter} ({changed} lignes calculees)"


def _apply_op_delete_column(sheet, op: dict) -> str:
    column = str(op.get("column") or "").strip().upper()
    if not column:
        raise ExcelToolError("delete_column requiert 'column'")
    sheet.delete_cols(column_index_from_string(column), 1)
    return f"colonne {column} supprimee"


def _apply_op_add_row(sheet, op: dict) -> str:
    values = op.get("values") or []
    at_row = op.get("atRow")
    if at_row:
        sheet.insert_rows(int(at_row))
        target_row = int(at_row)
    else:
        target_row = sheet.max_row + 1
    for i, value in enumerate(values, start=1):
        sheet.cell(row=target_row, column=i, value=value)
    return f"ligne ajoutee en {target_row} ({len(values)} valeurs)"


def _apply_op_delete_row(sheet, op: dict) -> str:
    row = op.get("row")
    if not row:
        raise ExcelToolError("delete_row requiert 'row'")
    sheet.delete_rows(int(row), 1)
    return f"ligne {row} supprimee"


def _apply_op_set_cell(sheet, op: dict) -> str:
    cell = str(op.get("cell") or "")
    col, row = _parse_cell_ref(cell)
    sheet.cell(row=row, column=col, value=op.get("value"))
    return f"cellule {cell} mise a jour"


def _apply_op_set_range(sheet, op: dict) -> str:
    range_ref = str(op.get("range") or "")
    values = op.get("values") or []
    start_col, start_row, _, _ = _parse_range_ref(range_ref)
    changed = 0
    for r, row_values in enumerate(values):
        for c, value in enumerate(row_values):
            sheet.cell(row=start_row + r, column=start_col + c, value=value)
            changed += 1
    return f"plage {range_ref} mise a jour ({changed} cellules)"


def _apply_op_write_formula(sheet, op: dict) -> str:
    range_ref = op.get("range")
    formula_template = str(op.get("formula") or "")
    if not formula_template:
        raise ExcelToolError("write_formula requiert 'formula'")
    if range_ref:
        start_col, start_row, end_col, end_row = _parse_range_ref(range_ref)
        changed = 0
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                sheet.cell(row=row, column=col, value=formula_template.replace("{row}", str(row)))
                changed += 1
        return f"formule ecrite sur {range_ref} ({changed} cellules)"
    cell = op.get("cell")
    if not cell:
        raise ExcelToolError("write_formula requiert 'cell' ou 'range'")
    col, row = _parse_cell_ref(cell)
    sheet.cell(row=row, column=col, value=formula_template.replace("{row}", str(row)))
    return f"formule ecrite en {cell}"


def _apply_op_copy_range(sheet, op: dict) -> str:
    source_range = str(op.get("sourceRange") or "")
    dest_cell = str(op.get("destCell") or "")
    start_col, start_row, end_col, end_row = _parse_range_ref(source_range)
    dest_col, dest_row = _parse_cell_ref(dest_cell)
    copied = 0
    for r in range(end_row - start_row + 1):
        for c in range(end_col - start_col + 1):
            value = sheet.cell(row=start_row + r, column=start_col + c).value
            sheet.cell(row=dest_row + r, column=dest_col + c, value=value)
            copied += 1
    return f"plage {source_range} copiee vers {dest_cell} ({copied} cellules)"


def _apply_op_clear_range(sheet, op: dict) -> str:
    range_ref = str(op.get("range") or "")
    start_col, start_row, end_col, end_row = _parse_range_ref(range_ref)
    cleared = 0
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            sheet.cell(row=row, column=col, value=None)
            cleared += 1
    return f"plage {range_ref} effacee ({cleared} cellules)"


def _apply_op_sort_range(sheet, op: dict) -> str:
    range_ref = str(op.get("range") or "")
    key_column = int(op.get("keyColumn") or 1)
    ascending = bool(op.get("ascending", True))
    has_header = bool(op.get("hasHeader", True))
    start_col, start_row, end_col, end_row = _parse_range_ref(range_ref)
    data_start = start_row + 1 if has_header else start_row
    rows = [
        [sheet.cell(row=r, column=c).value for c in range(start_col, end_col + 1)]
        for r in range(data_start, end_row + 1)
    ]
    rows.sort(key=lambda r: (r[key_column - 1] is None, r[key_column - 1]), reverse=not ascending)
    for i, row_values in enumerate(rows):
        for c, value in enumerate(row_values):
            sheet.cell(row=data_start + i, column=start_col + c, value=value)
    return f"plage {range_ref} triee sur la colonne {key_column} ({'croissant' if ascending else 'decroissant'})"


def _apply_op_filter_rows(sheet, op: dict) -> str:
    """Non destructif par design (mission §Original preservation, meme
    esprit applique ici) : ecrit les lignes correspondantes dans une NOUVELLE
    feuille plutot que de supprimer les lignes non correspondantes de la
    feuille source."""
    range_ref = str(op.get("range") or "")
    column = int(op.get("column") or 1)
    operator = str(op.get("operator") or "eq")
    value = op.get("value")
    has_header = bool(op.get("hasHeader", True))
    start_col, start_row, end_col, end_row = _parse_range_ref(range_ref)

    def matches(cell_value) -> bool:
        if operator == "eq":
            return str(cell_value) == str(value)
        if operator == "contains":
            return str(value).lower() in str(cell_value or "").lower()
        try:
            if operator == "gt":
                return float(cell_value) > float(value)
            if operator == "lt":
                return float(cell_value) < float(value)
        except (TypeError, ValueError):
            return False
        return False

    result_sheet = sheet.parent.create_sheet(title=f"{sheet.title}_filtre"[:31])
    out_row = 1
    if has_header:
        for c in range(start_col, end_col + 1):
            result_sheet.cell(row=1, column=c - start_col + 1, value=sheet.cell(row=start_row, column=c).value)
        out_row = 2
    data_start = start_row + 1 if has_header else start_row
    matched = 0
    for r in range(data_start, end_row + 1):
        cell_value = sheet.cell(row=r, column=column).value
        if matches(cell_value):
            for c in range(start_col, end_col + 1):
                result_sheet.cell(row=out_row, column=c - start_col + 1, value=sheet.cell(row=r, column=c).value)
            out_row += 1
            matched += 1
    return f"{matched} ligne(s) filtree(s) vers la nouvelle feuille '{result_sheet.title}'"


_OPERATION_HANDLERS = {
    "add_column": _apply_op_add_column,
    "delete_column": _apply_op_delete_column,
    "add_row": _apply_op_add_row,
    "delete_row": _apply_op_delete_row,
    "set_cell": _apply_op_set_cell,
    "set_range": _apply_op_set_range,
    "write_formula": _apply_op_write_formula,
    "copy_range": _apply_op_copy_range,
    "clear_range": _apply_op_clear_range,
    "sort_range": _apply_op_sort_range,
    "filter_rows": _apply_op_filter_rows,
}


def _apply_op_add_sheet(workbook, op: dict) -> str:
    name = str(op.get("name") or "").strip()[:31]
    if not name:
        raise ExcelToolError("add_sheet requiert 'name'")
    workbook.create_sheet(title=name)
    return f"feuille '{name}' ajoutee"


def _apply_op_rename_sheet(workbook, op: dict) -> str:
    old_name = str(op.get("oldName") or "")
    new_name = str(op.get("newName") or "").strip()[:31]
    if old_name not in workbook.sheetnames:
        raise ExcelToolError(f"feuille introuvable ({old_name})")
    if not new_name:
        raise ExcelToolError("rename_sheet requiert 'newName'")
    workbook[old_name].title = new_name
    return f"feuille '{old_name}' renommee en '{new_name}'"


def _apply_op_delete_sheet(workbook, op: dict) -> str:
    name = str(op.get("name") or "")
    if name not in workbook.sheetnames:
        raise ExcelToolError(f"feuille introuvable ({name})")
    if len(workbook.sheetnames) <= 1:
        raise ExcelToolError("impossible de supprimer la derniere feuille du classeur")
    del workbook[name]
    return f"feuille '{name}' supprimee"


_WORKBOOK_LEVEL_OPERATIONS = {
    "add_sheet": _apply_op_add_sheet,
    "rename_sheet": _apply_op_rename_sheet,
    "delete_sheet": _apply_op_delete_sheet,
}


def _preview_table(sheet, max_rows: int = 15) -> str:
    lines = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_rows), values_only=True):
        lines.append(" | ".join("" if v is None else str(v) for v in row))
    suffix = "" if sheet.max_row <= max_rows else f"\n... ({sheet.max_row - max_rows} lignes supplementaires non affichees)"
    return "\n".join(lines) + suffix


def apply_edit(file_id: str, user_id: str, sheet_name: str | None, operation: dict) -> dict:
    """Point d'entree unique (mission §5.1/§5.2). Verification de
    permission AVANT toute lecture/ecriture (reutilise
    workspace.user_can_access_file -- meme perimetre que les autres actions
    sur pieces jointes de cette app, aucun affaiblissement)."""
    if not workspace.user_can_access_file(file_id, user_id):
        raise ExcelToolError("acces refuse a ce fichier")

    op_type = str((operation or {}).get("type") or "")
    if op_type not in SUPPORTED_OPERATIONS:
        raise ExcelToolError(
            f"operation non supportee ({op_type}). Operations disponibles : {', '.join(sorted(SUPPORTED_OPERATIONS))}"
        )

    original_file = workspace.get_file(file_id)
    path = _storage_path(file_id)
    try:
        workbook = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        raise ExcelToolError(f"classeur illisible ({str(exc)[:150]})") from exc

    if op_type in _WORKBOOK_LEVEL_OPERATIONS:
        summary = _WORKBOOK_LEVEL_OPERATIONS[op_type](workbook, operation)
        preview_sheet = workbook.active
    else:
        sheet = _resolve_sheet(workbook, sheet_name)
        summary = _OPERATION_HANDLERS[op_type](sheet, operation)
        preview_sheet = sheet

    # Validation minimale (mission §Original preservation, "valider que le
    # classeur se rouvre") : on re-serialise en memoire et on la relit avant
    # d'ecrire le fichier final -- une corruption introduite par une
    # operation se traduit par une erreur explicite, jamais un fichier
    # livre silencieusement casse.
    import io

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    try:
        openpyxl.load_workbook(buffer)
    except Exception as exc:
        raise ExcelToolError(f"le classeur modifie ne peut pas etre rouvert ({str(exc)[:150]})") from exc

    new_file_id = f"file_{uuid.uuid4().hex}"
    base_name = (original_file or {}).get("name") or "classeur.xlsx"
    if "." in base_name:
        base_name = base_name.rsplit(".", 1)[0]
    new_name = f"{base_name} (modifie).xlsx"[:200]
    storage_name = f"{new_file_id}.xlsx"
    new_path = os.path.join(UPLOAD_DIR, storage_name)
    buffer.seek(0)
    content = buffer.read()
    with open(new_path, "wb") as handle:
        handle.write(content)

    workspace.create_file_record(
        file_id=new_file_id,
        user_id=user_id,
        original_name=new_name,
        mime_type=_XLSX_MIME,
        size_bytes=len(content),
        storage_reference=storage_name,
    )

    return {
        "newFileId": new_file_id,
        "newFileName": new_name,
        "sheet": preview_sheet.title,
        "summary": summary,
        "previewMarkdownTable": _preview_table(preview_sheet),
    }
